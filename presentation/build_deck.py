"""Build the NSP KPK Part III presentation (initial slides) on the PART III - PRESENTATION master.

Slides:
  1  Improving What Is In Place
  2  Nine aspects, from Afghanistan to Articles 232 and 234
  3  Afghanistan: four major takeaways (nine-aspect strip, Afghanistan highlighted)
  4  Takeaway 1: border trade gauge (native chart)
  5  Takeaway 2: people and access
  6  Takeaway 3: dialogue and verification
  7  Takeaway 4: connectivity and shared risk

Sources:
  - Conclusions (.md): unstruck text only (Improving section, Afghanistan Major Takeaways 1 to 4)
  - Abridged (.md): data notes 3 and 4 only, tied to the unstruck item 3c(11)

Run: python build_deck.py  (requires python-pptx)
"""

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "template" / "PART III - PRESENTATION.pptx"
OUT = HERE / "NSP_KPK_Part_III_Afghanistan_Draft.pptx"

# ---- Design tokens, taken from the PART III - PRESENTATION master ----
FONT = "Candara"                      # master title font; used for every run (one family)
DARK = RGBColor(0x00, 0x31, 0x00)     # master tab colour (003100)
GREEN = RGBColor(0x19, 0x6B, 0x24)    # theme accent 3, active highlight
GOLD = RGBColor(0xCC, 0x99, 0x00)     # master rule colour (CC9900)
TEXT = RGBColor(0x1A, 0x1A, 0x1A)
GREY_TXT = RGBColor(0x59, 0x59, 0x59)
LIGHT = RGBColor(0xF2, 0xF2, 0xF2)
MID = RGBColor(0xD0, 0xD0, 0xD0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

BODY = 15       # body text (pt)
HEAD = 18       # headings, maximum (pt)
STAT = 28       # data call-outs only (pt)
FOOT = 11       # source lines (pt)

SRC_CONCL = "Source: NSP KPK Part III Conclusions (unstruck text)"

ASPECTS = [
    ("Afghanistan", "The border can be a gauge instead of a switch"),
    ("Political", "The merger can be a contract instead of an event"),
    ("People", "Youth can be a workforce instead of a risk"),
    ("Civil Administration", "Posts can move forward before new ones are built"),
    ("Narrative and Information", "Pashto content can be a craft instead of a defence"),
    ("Technology, Intelligence and Information",
     "A connected population can be a channel instead of a target"),
    ("Economy", "Provincial resources can be a local stake instead of a dispute"),
    ("Military and LEAs", "Operations can end in convictions instead of clearances"),
    ("Articles 232 and 234", "The emergency can be a ladder instead of a pause"),
]
SHORT_NAMES = ["Afghanistan", "Political", "People", "Civil Admin", "Narrative",
               "Technology", "Economy", "Military and LEAs", "Articles 232 and 234"]

# Content area on the master: left of 1.0 in clears the NDU logo; right of 12.4 in matches the title.
X0, X1 = 1.0, 12.4
CW = X1 - X0

prs = Presentation(str(TEMPLATE))

# Remove the master's three sample slides; keep its masters, layouts and theme.
sld_ids = prs.slides._sldIdLst
for sld_id in list(sld_ids):
    prs.part.drop_rel(sld_id.rId)
    sld_ids.remove(sld_id)

LAYOUT = next(l for l in prs.slide_layouts if l.name == "Title Slide")


# ---------------- helpers ----------------
def _style_run(run, size, bold=False, color=TEXT):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _fill_text(tf, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    """paras: list of (text, size, bold, color) tuples."""
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, (txt, size, bold, color) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(4)
        r = p.add_run()
        r.text = txt
        _style_run(r, size, bold, color)


def box(slide, x, y, w, h, paras, fill=WHITE, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=0.14):
    """Shape with its text inside the shape. Solid fill keeps text readable over the master art."""
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.06
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.06)
    _fill_text(tf, paras, align=align, anchor=anchor)
    return s


def arrow(slide, x, y, w=0.4, h=0.32, fill=GOLD):
    s = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s


def text(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = Inches(0.02)
    _fill_text(tf, paras, align=align, anchor=anchor)
    return tb


def new_slide(title):
    """Slide on the master's Title Slide layout; the heading goes in the real title placeholder."""
    s = prs.slides.add_slide(LAYOUT)
    t = s.shapes.title
    t.text_frame.text = title
    for p in t.text_frame.paragraphs:
        for r in p.runs:
            r.font.name = FONT
            r.font.size = Pt(HEAD)
            r.font.bold = True
    return s


def footer(slide, source=SRC_CONCL):
    text(slide, X0, 6.62, CW, 0.3, [(source, FOOT, False, GREY_TXT)])


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body


def aspect_strip(slide, active=0, y=1.1, h=0.95):
    """Nine aspect tiles across the top, with the active aspect highlighted."""
    gap = 0.1
    w = (CW - gap * 8) / 9
    tiles = []
    for i, name in enumerate(SHORT_NAMES):
        on = i == active
        tiles.append(box(slide, X0 + i * (w + gap), y, w, h,
                         [(name, BODY, on, WHITE if on else GREY_TXT)],
                         fill=GREEN if on else LIGHT, align=PP_ALIGN.CENTER,
                         anchor=MSO_ANCHOR.MIDDLE, margin=0.04))
    return tiles


def add_click_steps(slide, steps):
    """Fade entrance, one On Click step per list element; shapes in a step appear together."""
    counter = [2]

    def nid():
        counter[0] += 1
        return counter[0]

    pars, bld = [], []
    for step in steps:
        outer, inner = nid(), nid()
        effects = []
        for k, shp in enumerate(step):
            sid = shp.shape_id
            node = "clickEffect" if k == 0 else "withEffect"
            e, s_, a = nid(), nid(), nid()
            effects.append(
                f'<p:par><p:cTn id="{e}" presetID="10" presetClass="entr" presetSubtype="0" '
                f'fill="hold" grpId="0" nodeType="{node}"><p:stCondLst><p:cond delay="0"/>'
                f'</p:stCondLst><p:childTnLst><p:set><p:cBhvr><p:cTn id="{s_}" dur="1" fill="hold">'
                f'<p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{sid}"/>'
                f'</p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>'
                f'</p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
                f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{a}" dur="400"/>'
                f'<p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
                f'</p:childTnLst></p:cTn></p:par>')
            if shp.has_chart:
                bld.append(f'<p:bldGraphic spid="{sid}" grpId="0"><p:bldAsOne/></p:bldGraphic>')
            else:
                bld.append(f'<p:bldP spid="{sid}" grpId="0" animBg="1"/>')
        pars.append(
            f'<p:par><p:cTn id="{outer}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
            f'</p:stCondLst><p:childTnLst><p:par><p:cTn id="{inner}" fill="hold"><p:stCondLst>'
            f'<p:cond delay="0"/></p:stCondLst><p:childTnLst>{"".join(effects)}</p:childTnLst>'
            f'</p:cTn></p:par></p:childTnLst></p:cTn></p:par>')

    xml = (
        f'<p:timing {nsdecls("p")}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" '
        f'nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" '
        f'dur="indefinite" nodeType="mainSeq"><p:childTnLst>{"".join(pars)}</p:childTnLst></p:cTn>'
        f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>'
        f'</p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/>'
        f'</p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
        f'<p:bldLst>{"".join(bld)}</p:bldLst></p:timing>')
    slide._element.append(parse_xml(xml))


# ---------------- slide 1: Improving What Is In Place ----------------
s = new_slide("Improving What Is In Place")
statement = text(s, X0, 1.1, CW, 0.9, [
    ("Most required actions are already in place, so the same effort can yield more with a change of direction.",
     BODY, True, TEXT),
    ("The position remains that a whole of nation effort is the most workable answer.", BODY, False, TEXT),
])
verbs = [("Flip", "Turn a threat's energy into an input"),
         ("Direct", "Steer existing activity towards a lawful channel"),
         ("Activate", "Use structures that exist on paper"),
         ("Learn", "Take a lesson from each result")]
vg = 0.4
vw = (CW - vg * 3) / 4
verb_shapes = []
for i, (name, desc) in enumerate(verbs):
    x = X0 + i * (vw + vg)
    on = i == 0
    verb_shapes.append(box(s, x, 2.25, vw, 1.75,
                           [(name, HEAD, True, WHITE if on else DARK),
                            (desc, BODY, False, WHITE if on else TEXT)],
                           fill=GREEN if on else WHITE, anchor=MSO_ANCHOR.TOP, margin=0.2))
    if i < 3:
        arrow(s, x + vw + 0.06, 2.95, w=0.28, h=0.3)
fw_label = text(s, X0, 4.3, CW, 0.4, [("The Framework: each point follows one construct", BODY, True, DARK)])
fw = [("Challenge", "What was treated as a problem"),
      ("What it can be", "The reframed possibility"),
      ("Opportunity", "What the reframe makes available")]
fw_w = (CW - 0.6 * 2) / 3
fw_boxes = []
for i, (name, desc) in enumerate(fw):
    x = X0 + i * (fw_w + 0.3)
    fw_boxes.append(box(s, x, 4.8, fw_w, 1.2, [(name, BODY, True, DARK), (desc, BODY, False, TEXT)],
                        fill=WHITE, line=GOLD))
    if i < 2:
        arrow(s, x + fw_w + 0.07, 5.25, w=0.26, h=0.3)
footer(s)
add_click_steps(s, [[statement], [verb_shapes[0]], [verb_shapes[1]], [verb_shapes[2]],
                    [verb_shapes[3]], [fw_label] + fw_boxes])
notes(s, "Improving what is in place. Four verbs steer the change: Flip, Direct, Activate, Learn. "
         "The Framework is applied to every field: challenge, what it can be, opportunity.")

# ---------------- slide 2: nine aspects ----------------
s = new_slide("Nine Aspects, From Afghanistan to Articles 232 and 234")
cg = 0.25
cw = (CW - cg * 2) / 3
ch, rg = 1.5, 0.2
card_by_row = [[], [], []]
for i, (name, desc) in enumerate(ASPECTS):
    r, c = divmod(i, 3)
    x = X0 + c * (cw + cg)
    y = 1.2 + r * (ch + rg)
    on = i == 0
    shp = box(s, x, y, cw, ch,
              [(f"{i + 1}  {name}", BODY, True, WHITE if on else DARK),
               (desc, BODY, False, WHITE if on else TEXT)],
              fill=GREEN if on else WHITE, anchor=MSO_ANCHOR.MIDDLE, margin=0.2)
    card_by_row[r].append(shp)
footer(s)
add_click_steps(s, [card_by_row[0], card_by_row[1], card_by_row[2]])
notes(s, "The discussion covers nine aspects, from Afghanistan to Articles 232 and 234. "
         "Afghanistan is the first and is taken up in the next slides. Each card's sub-line is the "
         "headline of that aspect in the Conclusions document.")

# ---------------- slide 3: Afghanistan overview ----------------
s = new_slide("Afghanistan: The Border Can Be a Gauge Instead of a Switch")
strip = aspect_strip(s, active=0)
takeaways = [
    ("Takeaway 1", "Border Trade Gauge",
     "The border has been run as a switch. Quota, conduct, surety, conditions, venue and rule can "
     "turn it into a gauge that answers to measured conduct."),
    ("Takeaway 2", "People and Access",
     "Kin and patients have been seen as a risk or a burden. A managed channel keeps kinship and "
     "care inside lawful and checked paths and earns goodwill."),
    ("Takeaway 3", "Dialogue and Verification",
     "Talks have rested on assurances. Custom, an independent verifier and area language give "
     "dialogue a structure that Pakistan can stand behind."),
    ("Takeaway 4", "Connectivity and Shared Risk",
     "Connectivity has been a reward after stability. Humanitarian data and a first segment inside "
     "Pakistan make it a shared interest that starts now."),
]
tg = 0.2
tw = (CW - tg * 3) / 4
t_cards = []
for i, (lab, name, body) in enumerate(takeaways):
    x = X0 + i * (tw + tg)
    t_cards.append(box(s, x, 2.45, tw, 3.9,
                       [(lab, BODY, True, GREEN), (name, HEAD, True, DARK), (body, BODY, False, TEXT)],
                       fill=WHITE, anchor=MSO_ANCHOR.TOP, margin=0.2))
footer(s)
add_click_steps(s, [strip] + [[c] for c in t_cards])
notes(s, "Afghanistan is the first of nine aspects. Four major takeaways follow, one per slide. "
         "Takeaway 1 is supported by figures; takeaways 2 to 4 are framed as shifts in how the issue is treated.")

# ---------------- slide 4: takeaway 1, border trade gauge ----------------
s = new_slide("Border Trade Gauge: The Border Has Been Run as a Switch")
strip = aspect_strip(s, active=0)
chart_title = text(s, X0, 2.45, 5.5, 0.4, [("Exports to Afghanistan, USD billion", BODY, True, DARK)])
cd = CategoryChartData()
cd.categories = ["FY2015", "FY2024-25"]
cd.add_series("Exports to Afghanistan", (1.699, 0.774))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(X0), Inches(2.9),
                        Inches(5.5), Inches(3.5), cd)
ch_ = gf.chart
ch_.has_legend = False
ch_.font.name = FONT
ch_.font.size = Pt(BODY)
ch_.value_axis.visible = False
ch_.value_axis.has_major_gridlines = False
ch_.category_axis.format.line.color.rgb = MID
plot = ch_.plots[0]
plot.gap_width = 90
plot.has_data_labels = True
dl = plot.data_labels
dl.number_format = '"$"0.00"bn"'
dl.number_format_is_linked = False
dl.position = XL_LABEL_POSITION.OUTSIDE_END
dl.font.size = Pt(BODY)
dl.font.bold = True
ser = plot.series[0]
ser.format.fill.solid()
ser.format.fill.fore_color.rgb = MID
pt = ser.points[1]
pt.format.fill.solid()
pt.format.fill.fore_color.rgb = GREEN

stat_a = box(s, 6.9, 2.45, 2.6, 1.6,
             [("-28%", STAT, True, GREEN),
              ("Bilateral trade, 2024 to 2025 (USD 2.46bn to 1.77bn)", BODY, False, TEXT)],
             fill=WHITE, anchor=MSO_ANCHOR.TOP, margin=0.18)
stat_b = box(s, 9.7, 2.45, 2.7, 1.6,
             [("-54%", STAT, True, GREEN),
              ("Exports to Afghanistan, FY2015 to FY2024-25", BODY, False, TEXT)],
             fill=WHITE, anchor=MSO_ANCHOR.TOP, margin=0.18)
switch = box(s, 6.9, 4.3, 2.6, 1.35,
             [("Switch", BODY, True, DARK), ("Open or shut, with nothing in between", BODY, False, TEXT)],
             fill=WHITE, line=MID)
arr = arrow(s, 9.55, 4.8, w=0.5, h=0.32, fill=GOLD)
gauge = box(s, 10.1, 4.3, 2.3, 1.35,
            [("Gauge", BODY, True, WHITE), ("Quota, conduct and published triggers", BODY, False, WHITE)],
            fill=GREEN)
policy = text(s, 6.9, 5.85, 5.5, 0.7, [
    ("Trade policy has swung between full closure and full opening, and each swing fed smugglers.",
     BODY, False, TEXT)])
footer(s, "Source: Part III Abridged, notes 3 and 4 (PIDE and PAJCCI estimates; State Bank of Pakistan export data)")
add_click_steps(s, [[chart_title, gf], [stat_a, stat_b], [switch, arr, gauge], [policy]])
notes(s, "Takeaway 1. Exports to Afghanistan fell from USD 1.70bn (FY2015) to USD 0.77bn (FY2024-25), "
         "a 54 per cent fall in ten years. Bilateral trade fell 28 per cent in 2025 on chamber and PIDE "
         "estimates. The switch-to-gauge point is the shift in policy the takeaway argues for.")

# ---------------- slide 5: takeaway 2, people and access ----------------
s = new_slide("People and Access: Kin and Patients Have Been Seen as a Risk or a Burden")
strip = aspect_strip(s, active=0)
left = box(s, X0, 2.6, 5.0, 2.9,
           [("Seen as", BODY, True, GREY_TXT),
            ("A risk or a burden", HEAD, True, DARK),
            ("Kin and patients have been seen as a risk or a burden.", BODY, False, TEXT)],
           fill=WHITE, anchor=MSO_ANCHOR.TOP, margin=0.25)
arr = arrow(s, 6.2, 3.9, w=1.0, h=0.42, fill=GOLD)
right = box(s, 7.4, 2.6, 5.0, 2.9,
            [("Can be", BODY, True, WHITE),
             ("A managed channel", HEAD, True, WHITE),
             ("Keeps kinship and care inside lawful and checked paths, and earns goodwill.",
              BODY, False, WHITE)],
            fill=GREEN, anchor=MSO_ANCHOR.TOP, margin=0.25)
bottom = text(s, X0, 5.85, CW, 0.6, [
    ("The opportunity: goodwill, earned through checked and lawful paths", BODY, True, DARK)])
footer(s)
add_click_steps(s, [[left], [arr, right], [bottom]])
notes(s, "Takeaway 2. Framed as a shift in treatment, not as data. The Conclusions text for this takeaway "
         "carries no figures in unstruck form, so none are shown here.")

# ---------------- slide 6: takeaway 3, dialogue and verification ----------------
s = new_slide("Dialogue and Verification: Talks Have Rested on Assurances")
strip = aspect_strip(s, active=0)
left = box(s, X0, 2.6, 3.4, 2.9,
           [("Talks have rested on", BODY, True, GREY_TXT),
            ("Assurances", HEAD, True, DARK)],
           fill=WHITE, anchor=MSO_ANCHOR.MIDDLE, margin=0.25)
arr_a = arrow(s, 4.6, 3.95, w=0.45, h=0.36, fill=GOLD)
inputs = []
for i, lab in enumerate(["Custom", "Independent verifier", "Area language"]):
    inputs.append(box(s, 5.25, 2.6 + i * 1.0, 3.6, 0.8, [(lab, BODY, True, DARK)],
                      fill=WHITE, line=GREEN))
arr_b = arrow(s, 9.0, 3.95, w=0.45, h=0.36, fill=GOLD)
right = box(s, 9.6, 2.6, 2.8, 2.9,
            [("A structure", BODY, True, WHITE),
             ("Pakistan can stand behind", HEAD, True, WHITE)],
            fill=GREEN, anchor=MSO_ANCHOR.MIDDLE, margin=0.25)
footer(s)
add_click_steps(s, [[left], inputs, [arr_a, arr_b, right]])
notes(s, "Takeaway 3. Three inputs turn assurances into a structure: custom, an independent verifier and "
         "area language. The Conclusions items that give examples for this takeaway are struck and are not shown.")

# ---------------- slide 7: takeaway 4, connectivity and shared risk ----------------
s = new_slide("Connectivity and Shared Risk: A Reward After Stability Can Become a Shared Interest")
strip = aspect_strip(s, active=0)
a = box(s, X0, 2.6, 3.6, 2.9,
        [("Today", BODY, True, GREY_TXT),
         ("Connectivity has been a reward after stability", BODY, False, TEXT)],
        fill=WHITE, anchor=MSO_ANCHOR.TOP, margin=0.25)
arr_a = arrow(s, 4.75, 3.95, w=0.45, h=0.36, fill=GOLD)
b = box(s, 5.3, 2.6, 3.6, 2.9,
        [("Start now", BODY, True, WHITE),
         ("Humanitarian data and a first segment inside Pakistan", BODY, False, WHITE)],
        fill=GREEN, anchor=MSO_ANCHOR.TOP, margin=0.25)
arr_b = arrow(s, 9.05, 3.95, w=0.45, h=0.36, fill=GOLD)
c = box(s, 9.6, 2.6, 2.8, 2.9,
        [("Result", BODY, True, DARK),
         ("A shared interest that starts now", HEAD, True, DARK)],
        fill=WHITE, line=GREEN, anchor=MSO_ANCHOR.TOP, margin=0.25)
footer(s)
add_click_steps(s, [[a], [arr_a, b], [arr_b, c]])
notes(s, "Takeaway 4. Connectivity moves from a reward after stability to a shared interest that starts now, "
         "through humanitarian data and a first segment built inside Pakistan.")

prs.save(OUT)
print("saved", OUT)
