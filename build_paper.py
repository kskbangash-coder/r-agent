# -*- coding: utf-8 -*-
"""
Render the research paper "Formulation of New Provinces in Pakistan:
Challenges and Prospects" to PDF (reportlab) and DOCX (python-docx).

Usage: python3 build_paper.py
Outputs: NewProvinces_Pakistan_Paper.pdf, NewProvinces_Pakistan_Paper.docx
"""
import re
import sys

from paper_content import TITLE, META_LINES, ABSTRACT, KEYWORDS, BLOCKS, REFERENCES

NAVY = (0x1F, 0x4E, 0x79)

# --------------------------------------------------------------------- markup
def xml_escape(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

def to_rl(t):
    """Inline markdown (**bold**, *italic*) -> reportlab paragraph markup."""
    t = xml_escape(t)
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    t = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"<i>\1</i>", t)
    return t

def to_runs(t):
    """Inline markdown -> list of (text, bold, italic) for python-docx."""
    out = []
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", t):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            out.append((part[2:-2], True, False))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append((part[1:-1], False, True))
        else:
            out.append((part, False, False))
    return out

# --------------------------------------------------------------------- PDF
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER, TA_LEFT
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph,
                                Spacer, Table, TableStyle, Image, KeepTogether)
from reportlab.pdfgen import canvas as _canvas

BODY = 9.8
LEAD = 13.6
NAVYC = colors.HexColor("#1F4E79")
GREY = colors.HexColor("#555555")
LT = colors.HexColor("#EEF2F7")


class NumberedCanvas(_canvas.Canvas):
    def __init__(self, *args, **kw):
        super().__init__(*args, **kw)
        self._saved = []

    def showPage(self):
        self._saved.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        n = len(self._saved)
        for st in self._saved:
            self.__dict__.update(st)
            self._draw_footer(n)
            super().showPage()
        super().save()

    def _draw_footer(self, total):
        W, H = A4
        self.saveState()
        self.setStrokeColor(colors.HexColor("#C9D3DD"))
        self.setLineWidth(0.4)
        self.line(doc.leftMargin, 11 * mm, W - doc.rightMargin, 11 * mm)
        self.setFont("Times-Roman", 7.2)
        self.setFillColor(GREY)
        self.drawString(doc.leftMargin, 8 * mm, "Formulation of New Provinces in Pakistan - Challenges and Prospects")
        self.drawRightString(W - doc.rightMargin, 8 * mm, "Page %d of %d" % (self._pageNumber, total))
        self.restoreState()


def pdf_styles():
    s = {}
    s["title"] = ParagraphStyle("t", fontName="Times-Bold", fontSize=16.5, leading=20,
                                textColor=NAVYC, alignment=TA_CENTER, spaceAfter=4)
    s["meta"] = ParagraphStyle("m", fontName="Times-Italic", fontSize=9, leading=12,
                               textColor=GREY, alignment=TA_CENTER, spaceAfter=1)
    s["absbox"] = ParagraphStyle("a", fontName="Times-Roman", fontSize=9.0, leading=12.6,
                                 alignment=TA_JUSTIFY, leftIndent=6, rightIndent=6,
                                 topPadding=5, bottomPadding=5)
    s["absh"] = ParagraphStyle("ah", fontName="Times-Bold", fontSize=9.6, leading=12,
                               textColor=NAVYC, spaceAfter=2)
    s["kw"] = ParagraphStyle("k", fontName="Times-Italic", fontSize=8.8, leading=11.6,
                              textColor=GREY, alignment=TA_JUSTIFY, spaceBefore=4, spaceAfter=6)
    s["h1"] = ParagraphStyle("h1", fontName="Times-Bold", fontSize=12.6, leading=15.4,
                              textColor=NAVYC, spaceBefore=13, spaceAfter=5, keepWithNext=1)
    s["h2"] = ParagraphStyle("h2", fontName="Times-Bold", fontSize=10.6, leading=13.2,
                              textColor=colors.HexColor("#2E5F8A"), spaceBefore=9, spaceAfter=3.4,
                              keepWithNext=1)
    s["p"] = ParagraphStyle("p", fontName="Times-Roman", fontSize=BODY, leading=LEAD,
                            alignment=TA_JUSTIFY, spaceAfter=5.2, firstLineIndent=15)
    s["p0"] = ParagraphStyle("p0", parent=s["p"], firstLineIndent=0)
    s["cap"] = ParagraphStyle("cap", fontName="Times-Italic", fontSize=8.4, leading=10.8,
                              textColor=GREY, alignment=TA_LEFT, spaceBefore=2, spaceAfter=6)
    s["tabh"] = ParagraphStyle("th", fontName="Times-Bold", fontSize=7.6, leading=9.4,
                               textColor=colors.white)
    s["cell"] = ParagraphStyle("c", fontName="Times-Roman", fontSize=7.6, leading=9.4)
    s["ref"] = ParagraphStyle("r", fontName="Times-Roman", fontSize=8.6, leading=11.2,
                              alignment=TA_JUSTIFY, leftIndent=15, firstLineIndent=-15,
                              spaceAfter=2.6)
    return s


def build_pdf(path):
    global doc
    s = pdf_styles()
    doc = BaseDocTemplate(path, pagesize=A4,
                          leftMargin=21 * mm, rightMargin=21 * mm,
                          topMargin=17 * mm, bottomMargin=18 * mm,
                          title=TITLE, subject=TITLE,
                          author="Arena research workflow")
    W, H = A4
    frame = Frame(doc.leftMargin, doc.bottomMargin,
                  W - doc.leftMargin - doc.rightMargin,
                  H - doc.topMargin - doc.bottomMargin, id="main",
                  leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    doc.addPageTemplates([PageTemplate(id="pt", frames=[frame])])
    avail = W - doc.leftMargin - doc.rightMargin

    story = [Paragraph(xml_escape(TITLE), s["title"])]
    for m in META_LINES:
        story.append(Paragraph(xml_escape(m), s["meta"]))
    story.append(Spacer(1, 6))

    ab = [[Paragraph("<b>Abstract.</b>  " + to_rl(ABSTRACT), s["absbox"])]]
    t = Table(ab, colWidths=[avail])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F6F8FB")),
        ("BOX", (0, 0), (-1, -1), 0.6, NAVYC),
    ]))
    story.append(t)
    story.append(Paragraph("<b>Keywords:</b> " + xml_escape(KEYWORDS), s["kw"]))

    prev = None
    for blk in BLOCKS:
        kind = blk[0]
        if kind == "h1":
            story.append(Paragraph(xml_escape(blk[1]), s["h1"]))
        elif kind == "h2":
            story.append(Paragraph(xml_escape(blk[1]), s["h2"]))
        elif kind == "p":
            style = s["p0"] if prev in (None, "h1", "h2") else s["p"]
            story.append(Paragraph(to_rl(blk[1]), style))
        elif kind == "fig":
            import os
            flow = []
            if os.path.exists(blk[1]):
                from reportlab.lib.utils import ImageReader
                iw, ih = ImageReader(blk[1]).getSize()
                w = avail * 0.92
                flow.append(Image(blk[1], width=w, height=w * ih / iw))
            flow.append(Paragraph(to_rl(blk[2]), s["cap"]))
            story.append(KeepTogether(flow))
        elif kind == "table":
            d = blk[1]
            header = [Paragraph(xml_escape(h), s["tabh"]) for h in d["head"]]
            data = [header]
            for r in d["rows"]:
                data.append([Paragraph(to_rl(c), s["cell"]) for c in r])
            colw = [f * avail for f in d["widths"]]
            t = Table(data, colWidths=colw, repeatRows=1)
            ts = [
                ("BACKGROUND", (0, 0), (-1, 0), NAVYC),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B9C6D4")),
                ("LEFTPADDING", (0, 0), (-1, -1), 3.4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3.4),
                ("TOPPADDING", (0, 0), (-1, -1), 2.4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2.6),
            ]
            for i in range(1, len(data)):
                if i % 2 == 0:
                    ts.append(("BACKGROUND", (0, i), (-1, i), LT))
            t.setStyle(TableStyle(ts))
            story.append(Spacer(1, 2))
            story.append(Paragraph("<b>" + xml_escape(d["title"]) + "</b>", s["cap"]))
            story.append(t)
            if d.get("note"):
                story.append(Paragraph(to_rl(d["note"]), s["cap"]))
        prev = kind

    story.append(Paragraph("References", s["h1"]))
    for r in REFERENCES:
        story.append(Paragraph(to_rl(r), s["ref"]))
    story.append(Spacer(1, 4))
    story.append(Paragraph(
        "<i>Note: prepared with the repository's AI research-agent workflow. Recent developments rely on "
        "official announcements and press reporting as cited; figures rounded. Status of proposals: "
        "mid-October 2026.</i>", s["cap"]))

    doc.build(story, canvasmaker=NumberedCanvas)
    print("PDF written:", path)


# --------------------------------------------------------------------- DOCX
def build_docx(path):
    from docx import Document
    from docx.shared import Pt, Mm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    import os

    doc = Document()
    for sec in doc.sections:
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin = sec.right_margin = Mm(24)
        sec.top_margin = sec.bottom_margin = Mm(20)

    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    st.paragraph_format.line_spacing = 1.18
    st.paragraph_format.space_after = Pt(6)

    def add_par(text, size=11, bold=False, italic=False, align="just", indent=None,
                space_before=None, space_after=None, color=None, hang=None):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.alignment = {"just": WD_ALIGN_PARAGRAPH.JUSTIFY, "c": WD_ALIGN_PARAGRAPH.CENTER,
                        "l": WD_ALIGN_PARAGRAPH.LEFT}[align]
        if indent: pf.first_line_indent = Mm(indent)
        if space_before is not None: pf.space_before = Pt(space_before)
        if space_after is not None: pf.space_after = Pt(space_after)
        if hang:
            pf.left_indent = Mm(hang); pf.first_line_indent = Mm(-hang)
        for (txt, b, i) in to_runs(text):
            r = p.add_run(txt)
            r.font.size = Pt(size)
            r.bold = bold or b
            r.italic = italic or i
            if color is not None:
                r.font.color.rgb = RGBColor(*color)
        return p

    def shade(cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)

    add_par(TITLE, size=16, bold=True, align="c", color=NAVY, space_after=2)
    for m in META_LINES:
        add_par(m, size=9.5, italic=True, align="c", color=(0x55, 0x55, 0x55), space_after=1)
    add_par("", size=4, space_after=2)

    ap = doc.add_paragraph()
    ap.paragraph_format.space_before = Pt(4)
    ap.paragraph_format.space_after = Pt(4)
    ap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    ap.paragraph_format.left_indent = Mm(6); ap.paragraph_format.right_indent = Mm(6)
    r = ap.add_run("Abstract.  ")
    r.bold = True; r.font.size = Pt(10); r.font.color.rgb = RGBColor(*NAVY)
    for (txt, b, i) in to_runs(ABSTRACT):
        rr = ap.add_run(txt)
        rr.font.size = Pt(10); rr.bold = b; rr.italic = i

    kp = add_par("Keywords: " + KEYWORDS, size=10, italic=True, color=(0x55, 0x55, 0x55))

    prev = None
    for blk in BLOCKS:
        kind = blk[0]
        if kind == "h1":
            add_par(blk[1], size=13, bold=True, align="l", color=NAVY,
                    space_before=14, space_after=5)
        elif kind == "h2":
            add_par(blk[1], size=11.5, bold=True, align="l", color=(0x2E, 0x5F, 0x8A),
                    space_before=9, space_after=3.5)
        elif kind == "p":
            add_par(blk[1], size=11, indent=None if prev in (None, "h1", "h2") else 5.5)
        elif kind == "fig":
            if os.path.exists(blk[1]):
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(blk[1], width=Mm(160))
            add_par(blk[2], size=9, italic=True, align="c", color=(0x55, 0x55, 0x55))
        elif kind == "table":
            d = blk[1]
            add_par(d["title"], size=9.5, bold=True, align="l", space_before=6, space_after=2)
            t = doc.add_table(rows=1 + len(d["rows"]), cols=len(d["head"]))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, h in enumerate(d["head"]):
                c = t.rows[0].cells[j]
                c.text = ""
                rr = c.paragraphs[0].add_run(h)
                rr.bold = True; rr.font.size = Pt(8.5)
                rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shade(c, "1F4E79")
            for i, row in enumerate(d["rows"]):
                for j, val in enumerate(row):
                    c = t.rows[i + 1].cells[j]
                    c.text = ""
                    parts = to_runs(val)
                    p0 = c.paragraphs[0]
                    for (txt, b, it) in parts:
                        rr = p0.add_run(txt)
                        rr.font.size = Pt(8.5); rr.bold = b; rr.italic = it
                    if i % 2 == 1:
                        shade(c, "EEF2F7")
            for j, f in enumerate(d["widths"]):
                wmm = int(162 * f)
                for row in t.rows:
                    row.cells[j].width = Mm(wmm)
            if d.get("note"):
                add_par(d["note"], size=8.5, italic=True, color=(0x55, 0x55, 0x55),
                        align="just", space_after=6)
        prev = kind

    add_par("References", size=13, bold=True, align="l", color=NAVY, space_before=14, space_after=5)
    for rref in REFERENCES:
        add_par(rref, size=9.5, align="just", hang=6.5, space_after=2.5)
    add_par("Note: prepared with the repository's AI research-agent workflow. Recent developments "
            "rely on official announcements and press reporting as cited; figures rounded. Status of "
            "proposals: mid-October 2026.", size=9, italic=True, color=(0x55, 0x55, 0x55))

    doc.save(path)
    print("DOCX written:", path)


if __name__ == "__main__":
    build_pdf("NewProvinces_Pakistan_Paper.pdf")
    build_docx("NewProvinces_Pakistan_Paper.docx")
